import io
import time
import zipfile

import pytest

from apps.presentations.generators import generate_presentations_for_context, generate_slide
from apps.presentations.pdf_export import generate_pdf_bytes
from apps.presentations.pptx_export import generate_pptx_bytes


@pytest.fixture(autouse=True)
def _jpeg_cache(tmp_path, monkeypatch):
    from apps.presentations import export_images
    monkeypatch.setattr(export_images, 'CACHE_DIR', tmp_path)
    export_images.load_jpeg.cache_clear()


def _decks(count=12):
    for meta in generate_presentations_for_context('export_test_ctx', count=count):
        slides = [generate_slide(meta, n) for n in range(1, meta['slide_count'] + 1)]
        yield meta, slides


def test_pptx_opens_in_python_pptx_with_all_slides_and_notes():
    from pptx import Presentation
    for meta, slides in _decks():
        prs = Presentation(io.BytesIO(generate_pptx_bytes(meta, slides)))
        assert len(prs.slides) == meta['slide_count']
        for slide, data in zip(prs.slides, slides):
            if data.get('speaker_note'):
                assert slide.notes_slide.notes_text_frame.text == data['speaker_note']


def test_pptx_carries_watermark_and_no_python_pptx_boilerplate():
    meta, slides = next(_decks(1))
    with zipfile.ZipFile(io.BytesIO(generate_pptx_bytes(meta, slides))) as z:
        core = z.read('docProps/core.xml').decode()
        assert meta['watermark_token'] in core
        assert 'python-pptx' not in core
        assert z.namelist()[0] == '[Content_Types].xml'
        # Keynote won't open a deck with notes unless the master is listed.
        assert '<p:notesMasterIdLst><p:notesMasterId r:id=' in z.read('ppt/presentation.xml').decode()


def test_pdf_is_well_formed_and_watermarked():
    for meta, slides in _decks(4):
        pdf = generate_pdf_bytes(meta, slides)
        assert pdf.startswith(b'%PDF-1.4')
        assert pdf.rstrip().endswith(b'%%EOF')
        assert pdf.count(b'/Type /Page ') == meta['slide_count']
        # xref offsets must point at the objects they index
        xref_pos = int(pdf.rsplit(b'startxref\n', 1)[1].split(b'\n')[0])
        entries = pdf[xref_pos:].split(b'\n')[3:]
        for n, entry in enumerate(entries[:10], 1):
            offset = int(entry[:10])
            assert pdf[offset:].startswith(f'{n} 0 obj'.encode())


def test_exports_are_fast():
    decks = list(_decks())
    for fn in (generate_pptx_bytes, generate_pdf_bytes):
        for meta, slides in decks:   # warm the JPEG cache
            fn(meta, slides)
        start = time.perf_counter()
        for meta, slides in decks:
            fn(meta, slides)
        per_deck = (time.perf_counter() - start) / len(decks)
        # WeasyPrint/python-pptx took ~3s / ~0.4s per deck; keep an order of
        # magnitude of headroom for slow CI while still catching a regression.
        assert per_deck < 0.1, f'{fn.__name__}: {per_deck:.3f}s per deck'
